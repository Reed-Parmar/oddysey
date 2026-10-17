"""Orchestration: entities -> scenarios -> engine labels -> files -> manifest."""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from pydantic import ValidationError

from scripts.synthetic import GENERATOR_VERSION, ingest, llm
from scripts.synthetic import submissions as subs_mod
from scripts.synthetic.allocation import ALLOCATION_POLICY
from scripts.synthetic.entities import SynthProject, SynthTrainee, make_projects, make_trainees
from scripts.synthetic.reference import Reference, load_reference
from scripts.synthetic.scenarios import (
    SCENARIO_TYPES,
    SPLITS,
    assign_entity_splits,
    build_scenario,
    jl,
    run_scenario,
    scenario_rows,
    scenario_splits,
    scenario_summary,
)
from scripts.synthetic.schemas import SCHEMA_FILES, TABLE_MODELS
from spa_ai.scoring import SCORING_VERSION

MANIFEST_KIND = "spa-synthetic-dataset"
FORMATS = ("csv", "jsonl", "xlsx", "import")
LIMITS = {"trainees": (30, 20000), "projects": (9, 2000), "scenarios": (3, 500)}


class GenerationError(Exception):
    """Invalid parameters, unusable inputs, an unsafe output directory or an inconsistent example."""


@dataclass
class Config:
    trainees: int = 1000
    projects: int = 100
    scenarios: int = 30
    seed: int = 42
    out: Path = Path("data/synthetic")
    formats: tuple[str, ...] = ("csv", "jsonl")
    submissions: bool = True
    llm: bool = True
    workbook: Path = Path("CSR_Hackathon_ITI_Project_Dataset.xlsx")
    trainee_csv: Path | None = None
    force: bool = False
    quiet: bool = False
    extra: dict[str, Any] = field(default_factory=dict)

    def check(self) -> None:
        for key, (lo, hi) in LIMITS.items():
            v = getattr(self, key)
            if not isinstance(v, int) or not lo <= v <= hi:
                raise GenerationError(f"--{key} must be an integer between {lo} and {hi} (got {v})")
        unknown = set(self.formats) - set(FORMATS)
        if unknown or "csv" not in self.formats:
            raise GenerationError(f"--formats must include csv and only use {', '.join(FORMATS)} (got {self.formats})")
        if "jsonl" not in self.formats and self.llm:
            self.llm = False


# ---------------------------------------------------------------- output directory safety


def prepare_out(cfg: Config) -> Path:
    out = cfg.out.resolve()
    sources = [cfg.workbook.resolve()] + ([cfg.trainee_csv.resolve()] if cfg.trainee_csv else [])
    for s in sources:
        if out == s or out in s.parents:
            raise GenerationError(f"output directory {out} contains a source file ({s.name}); choose another")
    if out.exists() and not out.is_dir():
        raise GenerationError(f"output path {out} is a file")
    if out.exists() and any(out.iterdir()):
        manifest = out / "manifest.json"
        try:
            previous = json.loads(manifest.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            previous = None
        if not previous or previous.get("kind") != MANIFEST_KIND:
            raise GenerationError(f"{out} is not empty and was not created by this generator; refusing to write")
        if not cfg.force:
            raise GenerationError(f"{out} already holds a generated dataset; pass --force to replace it")
        for rel in [*previous.get("files", {}), "manifest.json", "validation_report.json", "README.md"]:
            target = (out / rel).resolve()
            if out in target.parents and target.is_file():
                target.unlink()
        for sub in ("llm", "import", "ingest/scenarios", "ingest", "schema"):
            d = out / sub
            if d.is_dir() and not any(p.is_file() for p in d.rglob("*")):
                shutil.rmtree(d)
    out.mkdir(parents=True, exist_ok=True)
    return out


# ---------------------------------------------------------------- writers


def _cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    return v


# Header for tables that can legitimately be empty (e.g. a run with no rejections). write_csv checks every
# non-empty row against these lists, so they cannot drift from the row builders unnoticed.
COLUMNS: dict[str, list[str]] = {
    "allocations": [
        "scenario_id",
        "allocation_id",
        "trainee_id",
        "project_id",
        "project_code",
        "status",
        "source",
        "assignment_kind",
        "trigger",
        "created_at",
        "event_seq",
        "match_score",
        "match_mode",
        "factors",
        "explanation",
        "replaces_rejection_id",
        "replacement_rank",
        "scoring_version",
        "allocation_policy",
        "synthetic_record",
    ],
    "rejections": [
        "scenario_id",
        "rejection_id",
        "allocation_id",
        "trainee_id",
        "project_id",
        "project_code",
        "reason",
        "rejected_at",
        "event_seq",
        "previous_match_score",
        "previous_match_mode",
        "replacement_status",
        "replacement_attempts",
        "replacement_allocation_id",
        "selected_trainee_id",
        "selected_rank",
        "ranking_mode",
        "candidate_count",
        "eligible_count",
        "ranking_snapshot",
        "explanation",
        "synthetic_record",
    ],
    "recommendations": [
        "scenario_id",
        "split",
        "trainee_id",
        "rank",
        "project_id",
        "project_code",
        "score",
        "score_mode",
        "eligible",
        "recommended",
        "ineligible_reasons",
        "factors",
        "label_source",
        "scoring_version",
        "synthetic_record",
    ],
    "submissions": [
        "submission_id",
        "scenario_id",
        "allocation_id",
        "trainee_id",
        "project_id",
        "version",
        "submitted_at",
        "report_url",
        "video_url",
        "logbook_hours",
        "bom_cost_inr",
        "status",
        "synthetic_record",
    ],
    "evaluations": [
        "evaluation_id",
        "submission_id",
        "scenario_id",
        "evaluated_at",
        "outcome",
        "score_workmanship",
        "score_functional_demo",
        "score_safety_sop",
        "score_documentation",
        "score_viva",
        "total_score",
        "grade",
        "grade_label",
        "feedback",
        "rubric",
        "synthetic_record",
    ],
    **{name: list(model.model_fields) for name, model in TABLE_MODELS.items()},
}


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str] | None = None) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = columns or (list(rows[0]) if rows else [])
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(cols)
        for r in rows:
            if list(r) != cols:
                raise GenerationError(f"{path.name}: inconsistent columns")
            w.writerow([_cell(r[c]) for c in cols])
    return len(rows)


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False, separators=(",", ":")) + "\n")
    return len(rows)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------- entity rows


def trainee_row(t: SynthTrainee) -> dict[str, Any]:
    trades = t.trades
    return {
        "trainee_id": t.trainee_id,
        "full_name": t.full_name,
        "gender": t.gender,
        "centre_id": t.centre_id,
        "centre_name": t.centre_name,
        "primary_trade_id": trades[0].id,
        "primary_trade_name": trades[0].name,
        "trade_ids": jl([x.id for x in trades]),
        "trade_names": jl([x.name for x in trades]),
        "theory_score": t.theory,
        "practical_score": t.practical,
        "safety_score": t.safety,
        "assessment_status": "complete" if t.assessment_complete else "pending",
        "marks_on_record": t.marks_on_record,
        "skills": jl(t.skills),
        "skill_evidence": jl(t.skill_evidence),
        "career_interest": t.career_interest,
        "background": t.background,
        "education": t.education,
        "profile_archetype": t.archetype,
        "home_split": t.home_split,
        "synthetic_record": True,
    }


def project_row(p: SynthProject) -> dict[str, Any]:
    trades = p.trades
    return {
        "project_id": p.project_id,
        "project_code": p.project_code,
        "title": p.title,
        "description": p.description,
        "category": p.category,
        "domain": p.domain,
        "target_trade_id": trades[0].id,
        "target_trade_name": trades[0].name,
        "target_trade_ids": jl([x.id for x in trades]),
        "target_trade_names": jl([x.name for x in trades]),
        "difficulty": p.difficulty,
        "required_skills": jl(p.required_skills),
        "min_practical_score": p.min_practical_score,
        "min_match_score": p.min_match_score,
        "legacy_max_trainees": p.legacy_max_trainees,
        "catalog_status": p.catalog_status,
        "problem_doc_url": p.problem_doc_url,
        "requirements_version": p.requirements_version,
        "home_split": p.home_split,
        "synthetic_record": True,
    }


def write_raw_csv(path: Path, rows: list[list[Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        for r in rows:
            w.writerow([_cell(v) for v in r])


def write_xlsx(path: Path, tables: dict[str, list[dict[str, Any]]]) -> None:
    import openpyxl

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("README")
    ws.append(
        [
            "Synthetic dataset for Smart Project Allocation. Fictional records only. "
            "The CSV/JSONL files are canonical; this workbook is for inspection."
        ]
    )
    for name in ("scenarios", "trainees", "projects", "allocation_results", "rejections"):
        rows = tables.get(name) or []
        ws = wb.create_sheet(name[:31])
        if rows:
            ws.append(list(rows[0]))
            for r in rows:
                ws.append([_cell(v) for v in r.values()])
    wb.save(path)


# ---------------------------------------------------------------- main entry


def log(cfg: Config, msg: str) -> None:
    if not cfg.quiet:
        print(msg, flush=True)


def generate(cfg: Config) -> dict[str, Any]:
    cfg.check()
    try:
        ref: Reference = load_reference(cfg.workbook, cfg.trainee_csv)
    except Exception as e:  # ReferenceError and parser failures
        raise GenerationError(str(e)) from None
    out = prepare_out(cfg)
    for w in ref.warnings:
        log(cfg, f"warning: {w}")

    trainees = make_trainees(cfg.seed, cfg.trainees, ref)
    projects = make_projects(cfg.seed, cfg.projects)
    assign_entity_splits(cfg.seed, trainees, projects)
    t_by_id = {t.trainee_id: t for t in trainees}
    p_by_id = {p.project_id: p for p in projects}
    log(cfg, f"entities: {len(trainees)} trainees, {len(projects)} projects")

    splits = scenario_splits(cfg.seed, cfg.scenarios)
    scenarios = []
    tables: dict[str, list[dict[str, Any]]] = {
        k: []
        for k in (
            "candidate_pairs",
            "recommendations",
            "allocation_results",
            "allocations",
            "rejections",
            "scenario_projects",
        )
    }
    examples: dict[str, dict[str, list[dict[str, Any]]]] = {
        task: {s: [] for s in SPLITS} for task in ("rank_projects", "explain_allocation", "decision_output")
    }
    for i in range(cfg.scenarios):
        sc = build_scenario(cfg.seed, i, splits[i], trainees, projects)
        run_scenario(cfg.seed, sc, t_by_id)
        rows = scenario_rows(sc, t_by_id, p_by_id)
        for k, v in rows.items():
            tables[k].extend(v)
        if cfg.llm:
            results = {r["trainee_id"]: r for r in rows["allocation_results"]}
            for n, tid in enumerate(sorted(results), start=1):
                t = t_by_id[tid]
                try:
                    if any(p.allocatable for p in sc.projects):
                        examples["rank_projects"][sc.split].append(llm.rank_example(cfg.seed, sc, t, p_by_id, n))
                    examples["explain_allocation"][sc.split].append(
                        llm.allocation_example(sc, t, results[tid], p_by_id, n)
                    )
                    examples["decision_output"][sc.split].append(llm.decision_example(sc, t, results[tid], n))
                except (llm.ExampleError, ValidationError) as e:
                    raise GenerationError(f"{sc.scenario_id}/{tid}: rejected example: {e}") from None
        scenarios.append(sc)
        log(
            cfg,
            f"{sc.scenario_id} [{sc.split:<10}] {sc.type:<24} trainees={len(sc.cohort):>4} "
            f"projects={len(sc.projects):>3} pairs={len(rows['candidate_pairs']):>6}",
        )

    tables["scenarios"] = [scenario_summary(sc) for sc in scenarios]
    tables["trainees"] = [trainee_row(t) for t in trainees]
    tables["projects"] = [project_row(p) for p in projects]
    if cfg.submissions:
        tables["submissions"], tables["evaluations"] = subs_mod.build(cfg.seed, scenarios, t_by_id, p_by_id)

    for name, model in TABLE_MODELS.items():  # explicit schema validation before anything is written
        for r in tables[name]:
            try:
                model.model_validate(r)
            except ValidationError as e:
                raise GenerationError(
                    f"{name}: invalid row {r.get('trainee_id') or r.get('project_id')}: {e}"
                ) from None

    files: dict[str, dict[str, Any]] = {}

    def record(rel: str, rows: int) -> None:
        files[rel] = {"rows": rows}

    order = [
        "trainees",
        "projects",
        "scenarios",
        "scenario_projects",
        "candidate_pairs",
        "recommendations",
        "allocations",
        "rejections",
        "allocation_results",
        "submissions",
        "evaluations",
    ]
    for name in order:
        if name in tables:
            record(f"{name}.csv", write_csv(out / f"{name}.csv", tables[name], COLUMNS.get(name)))
    for name, model in TABLE_MODELS.items():
        p = out / "schema" / SCHEMA_FILES[name]
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(json.dumps(model.model_json_schema(), indent=2) + "\n", encoding="utf-8", newline="\n")
        record(f"schema/{SCHEMA_FILES[name]}", 0)
    if cfg.llm:
        for task, by_split in examples.items():
            for task_model, kind in ((llm.INPUT_MODELS[task], "input"), (llm.TARGET_MODELS[task], "target")):
                p = out / "schema" / f"llm_{task}.{kind}.schema.json"
                p.write_text(
                    json.dumps(task_model.model_json_schema(), indent=2) + "\n", encoding="utf-8", newline="\n"
                )
                record(f"schema/llm_{task}.{kind}.schema.json", 0)
            for split, split_rows in by_split.items():
                record(f"llm/{task}/{split}.jsonl", write_jsonl(out / "llm" / task / f"{split}.jsonl", split_rows))
    # Importer-ready files (always written): what an admin uploads. See ingest/README.md.
    ingest_files: dict[str, list[list[Any]]] = {}
    sups = ingest.supervisors(cfg.seed)
    ingest_files.update(ingest.chunks("ingest/supervisors", ingest.STAFF_HEADER, ingest.staff_rows(sups)))
    ingest_files.update(ingest.chunks("ingest/admins", ingest.STAFF_HEADER, ingest.staff_rows(ingest.admins(cfg.seed))))
    ingest_files.update(ingest.chunks("ingest/trainees", ingest.TRAINEE_HEADER, ingest.trainee_rows(trainees)))
    ingest_files.update(ingest.chunks("ingest/projects", ingest.PROJECT_HEADER, ingest.project_rows(projects)))
    if "import" in cfg.formats:
        for sc in scenarios:
            cohort = [t_by_id[tid] for tid, _ in sc.cohort]
            sproj = [p_by_id[p.project_id] for p in sc.projects]
            ingest_files.update(
                ingest.chunks(
                    f"ingest/scenarios/{sc.scenario_id}_trainees", ingest.TRAINEE_HEADER, ingest.trainee_rows(cohort)
                )
            )
            ingest_files.update(
                ingest.chunks(
                    f"ingest/scenarios/{sc.scenario_id}_projects", ingest.PROJECT_HEADER, ingest.project_rows(sproj)
                )
            )
    for rel, sheet in ingest_files.items():
        write_raw_csv(out / rel, sheet)
        record(rel, len(sheet) - 1)
    (out / "ingest" / "README.md").write_text(ingest.README, encoding="utf-8", newline="\n")
    record("ingest/README.md", 0)
    if "xlsx" in cfg.formats:
        write_xlsx(out / "inspection.xlsx", tables)
        record("inspection.xlsx", 0)  # zip timestamps make the bytes vary; excluded from reproducibility hashes

    for rel, meta in files.items():
        if rel != "inspection.xlsx":
            meta["sha256"] = sha256(out / rel)
    split_manifest = {
        "strategy": "entity-disjoint: trainees and projects are partitioned 70/15/15 by primary trade; each "
        "scenario samples only from its split's pools, so no scenario, trainee or project crosses splits",
        "scenarios": {sc.scenario_id: sc.split for sc in scenarios},
        "trainees": {s: sum(1 for t in trainees if t.home_split == s) for s in SPLITS},
        "projects": {s: sum(1 for p in projects if p.home_split == s) for s in SPLITS},
    }
    params = asdict(cfg)
    params = {
        k: (str(v) if isinstance(v, Path) else v) for k, v in params.items() if k not in ("quiet", "force", "extra")
    }
    params["workbook"] = Path(params["workbook"]).name
    params["trainee_csv"] = Path(params["trainee_csv"]).name if params["trainee_csv"] else None
    params["out"] = "(output directory)"
    manifest = {
        "kind": MANIFEST_KIND,
        "generator_version": GENERATOR_VERSION,
        "scoring_version": SCORING_VERSION,
        "allocation_policy": ALLOCATION_POLICY,
        "seed": cfg.seed,
        "parameters": params,
        "synthetic": True,
        "sources": [Path(s).name for s in ref.sources],
        "reference_stats": ref.stats(),
        "scenario_types": [t for t, _ in SCENARIO_TYPES],
        "splits": split_manifest,
        "files": files,
    }
    (out / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n"
    )
    return manifest
