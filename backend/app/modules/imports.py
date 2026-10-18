"""IMP-01 (trainees) and IMP-02 (projects): XLSX/CSV in the dataset's format (api.md C-11).

All-or-nothing: rows are validated and upserted inside one transaction, which is rolled back for a
dry run or when any row fails. Dry runs therefore report exactly what an apply would do."""

import csv
import io
import re
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Annotated, Any

import openpyxl
from fastapi import APIRouter, File, Form, UploadFile
from pydantic import BaseModel
from sqlalchemy import func, select

from app import audit, jobs
from app.core.config import get_settings
from app.core.deps import DB, Admin
from app.core.errors import AppError, field_error, invalid_state
from app.core.responses import Envelope
from app.db import models as M
from app.db.enums import (
    AllocationStatus,
    CycleStatus,
    Difficulty,
    ProjectProcessing,
    ProjectStatus,
    UserRole,
    UserStatus,
)
from app.domain import (
    clean_names,
    lock_cycle,
    pick_skills,
    project_input,
    resolve_skills_bulk,
    resolve_trade_names,
    set_project_trades,
    set_trainee_trades,
    trainee_input,
)
from spa_ai.scoring import eligibility

router = APIRouter(tags=["imports"])

MAX_ROWS = 2000


class RowError(BaseModel):
    row: int | None
    column: str | None
    message: str


class ImportResult(BaseModel):
    applied: bool
    rows: int
    created: int
    updated: int
    unchanged: int
    created_trades: list[str]
    created_skills: list[str]
    warnings: list[str]
    errors: list[RowError]


# ---------------------------------------------------------------- parsing


@dataclass
class Sheet:
    rows: list[tuple[int, dict[str, Any]]]  # (1-based sheet row number, {normalised header: value})
    headers: dict[str, str]  # normalised → original
    errors: list[RowError] = field(default_factory=list)


def _norm(h: object) -> str:
    return " ".join(str(h).split()).lower() if h is not None else ""


def _table(raw_rows: list[tuple[int, list[Any]]], marker: str, required: list[str]) -> Sheet:
    header_idx = next(
        (i for i, (_, cells) in enumerate(raw_rows) if any(_norm(c) == marker.lower() for c in cells)), None
    )
    if header_idx is None:
        return Sheet([], {}, [RowError(row=None, column=marker, message=f"Header row with '{marker}' not found")])
    _, header_cells = raw_rows[header_idx]
    headers = {_norm(c): " ".join(str(c).split()) for c in header_cells if _norm(c)}
    keys = [_norm(c) for c in header_cells]
    errors = [
        RowError(row=None, column=c, message="Required column is missing") for c in required if c.lower() not in headers
    ]
    rows: list[tuple[int, dict[str, Any]]] = []
    for number, cells in raw_rows[header_idx + 1 :]:
        if all(c is None or str(c).strip() == "" for c in cells):
            continue
        rows.append((number, {k: (cells[i] if i < len(cells) else None) for i, k in enumerate(keys) if k}))
    if len(rows) > MAX_ROWS:
        errors.append(RowError(row=None, column=None, message=f"At most {MAX_ROWS} data rows are allowed"))
    return Sheet(rows, headers, errors)


def parse_upload(filename: str, data: bytes, sheet_name: str, marker: str, required: list[str]) -> Sheet:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "xlsx" and data[:4] == b"PK\x03\x04":
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise AppError(415, "UNSUPPORTED_FILE_TYPE", "The file is not a valid .xlsx workbook") from None
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
        raw = [(i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), start=1)]
        wb.close()
        return _table(raw, marker, required)
    if ext == "csv" and b"\x00" not in data[:4096]:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise AppError(415, "UNSUPPORTED_FILE_TYPE", "CSV files must be UTF-8 encoded") from None
        raw = [(i, list(r)) for i, r in enumerate(csv.reader(io.StringIO(text)), start=1)]
        return _table(raw, marker, required)
    raise AppError(415, "UNSUPPORTED_FILE_TYPE", "Upload an .xlsx or .csv file")


def _text(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = " ".join(str(v).split())
    return s or None


def _decimal(v: Any) -> Decimal | None:
    s = _text(v)
    if s is None:
        return None
    try:
        return Decimal(s.rstrip("%"))
    except InvalidOperation:
        raise ValueError("must be a number") from None


def _split(v: Any) -> list[str]:
    s = _text(v)
    return [p.strip() for p in s.split(",") if p.strip()] if s else []


def same_trades(primary_id: Any, current: list[M.Trade], row: list[M.Trade]) -> bool:
    """Trades are a set plus a primary (first) trade. The loaded relationship is ordered by name, so
    comparing ordered lists flagged every multi-trade row as changed on each re-import."""
    return bool(row) and primary_id == row[0].id and {t.id for t in current} == {t.id for t in row}


CODE_RE = re.compile(r"^[A-Za-z0-9-]{1,40}$")


class RowCollector:
    def __init__(self, sheet: Sheet) -> None:
        self.errors = list(sheet.errors)
        self.headers = sheet.headers

    def err(self, row: int, column: str, message: str) -> None:
        self.errors.append(RowError(row=row, column=self.headers.get(column.lower(), column), message=message))


# ---------------------------------------------------------------- shared endpoint plumbing


async def _read(file: UploadFile) -> bytes:
    limit = get_settings().max_import_bytes
    data = await file.read(limit + 1)
    if len(data) > limit:
        raise AppError(413, "FILE_TOO_LARGE", "Import files must be 5 MB or smaller")
    return data


def _cycle(db: DB, cycle_id: str) -> M.Cycle:
    try:
        cid = uuid.UUID(cycle_id)
    except ValueError:
        raise field_error("body.cycle_id", "must be a UUID") from None
    if db.get(M.Cycle, cid) is None:
        raise field_error("body.cycle_id", "unknown cycle")
    cycle = lock_cycle(db, cid)
    if cycle.status == CycleStatus.archived:
        raise invalid_state("Cycle is archived (read-only)")
    return cycle


def _finish(db: DB, result: ImportResult, dry_run: bool, admin: M.User, action: str, cycle: M.Cycle) -> ImportResult:
    if result.errors:
        db.rollback()
        if not dry_run:
            raise AppError(
                422,
                "IMPORT_INVALID",
                "The file has errors; nothing was imported",
                [
                    {"field": f"row {e.row}: {e.column}" if e.row else (e.column or "file"), "message": e.message}
                    for e in result.errors
                ],
            )
        return result
    if dry_run:
        db.rollback()
        return result
    audit.record(
        db,
        admin.id,
        action,
        "cycle",
        cycle.id,
        {"rows": result.rows, "created": result.created, "updated": result.updated},
    )
    db.commit()
    result.applied = True
    return result


# ---------------------------------------------------------------- IMP-01 trainees

TRAINEE_REQUIRED = [
    "Trainee ID",
    "Full Name",
    "Gender",
    "ITI Trade",
    "ITI Name",
    "Trade Theory Score (%)",
    "Workshop Practical Score (%)",
    "Safety Compliance Score (%)",
    "Core Skills Demonstrated",
    "Primary Career Interest",
]
GENDERS = {"male": "male", "female": "female", "other": "other", "undisclosed": "undisclosed"}


@router.post("/admin/imports/trainees")
async def import_trainees(
    admin: Admin,
    db: DB,
    file: Annotated[UploadFile, File()],
    cycle_id: Annotated[str, Form()],
    dry_run: Annotated[bool, Form()] = True,
) -> Envelope[ImportResult]:
    data = await _read(file)
    sheet = parse_upload(file.filename or "", data, "Trainee Profiles", "Trainee ID", TRAINEE_REQUIRED)
    cycle = _cycle(db, cycle_id)
    rc = RowCollector(sheet)
    result = ImportResult(
        applied=False,
        rows=len(sheet.rows),
        created=0,
        updated=0,
        unchanged=0,
        created_trades=[],
        created_skills=[],
        warnings=[],
        errors=[],
    )
    if rc.errors:
        result.errors = rc.errors
        return Envelope(data=_finish(db, result, dry_run, admin, "import.trainees", cycle))

    seen: set[str] = set()
    parsed: list[tuple[int, dict[str, Any]]] = []
    for number, r in sheet.rows:
        before = len(rc.errors)
        code = _text(r.get("trainee id"))
        if not code or not CODE_RE.match(code):
            rc.err(number, "Trainee ID", "required; letters, digits and '-' only (max 40)")
        elif code.upper() in seen:
            rc.err(number, "Trainee ID", f"duplicate code {code} in this file")
        else:
            seen.add(code.upper())
        name = _text(r.get("full name"))
        if not name or len(name) > 120:
            rc.err(number, "Full Name", "required (1–120 characters)")
        gender_raw = _text(r.get("gender"))
        gender = None
        if gender_raw:
            gender = GENDERS.get(gender_raw.lower())
            if gender is None:
                rc.err(number, "Gender", "must be Male, Female, Other or blank")
        trades_raw = _split(r.get("iti trade"))
        if not trades_raw or len(trades_raw) > 5 or any(len(x) > 80 for x in trades_raw):
            rc.err(number, "ITI Trade", "required: 1–5 trades separated by commas (each up to 80 characters)")
        centre = _text(r.get("iti name"))
        if centre and len(centre) > 120:
            rc.err(number, "ITI Name", "at most 120 characters")
        scores: dict[str, Decimal | None] = {}
        for col, key in (
            ("Trade Theory Score (%)", "theory_score"),
            ("Workshop Practical Score (%)", "practical_score"),
            ("Safety Compliance Score (%)", "safety_score"),
        ):
            try:
                v = _decimal(r.get(col.lower()))
                if v is not None and not (0 <= v <= 100):
                    raise ValueError("must be 0–100")
                scores[key] = v
            except ValueError as e:
                rc.err(number, col, str(e))
        skills = _split(r.get("core skills demonstrated"))
        if len(skills) > 20 or any(len(s) > 80 for s in skills):
            rc.err(number, "Core Skills Demonstrated", "at most 20 skills, each up to 80 characters")
        interest = _text(r.get("primary career interest"))
        if interest and len(interest) > 300:
            rc.err(number, "Primary Career Interest", "at most 300 characters")
        background = _text(r.get("background")) if "background" in sheet.headers else None
        if background and len(background) > 2000:
            rc.err(number, "Background", "at most 2000 characters")
        if len(rc.errors) == before:
            parsed.append(
                (
                    number,
                    {
                        "code": code,
                        "full_name": name,
                        "gender": gender,
                        "trades": clean_names(trades_raw, 5, "ITI Trade") if trades_raw else [],
                        "centre_name": centre,
                        **scores,
                        "skills": skills,
                        "career_interest": interest,
                        "background": background,
                        "has_background": "background" in sheet.headers,
                    },
                )
            )
    result.errors = rc.errors
    if result.errors:
        return Envelope(data=_finish(db, result, dry_run, admin, "import.trainees", cycle))

    trades, created_trades = resolve_trade_names(
        db, sorted({n for _, p in parsed for n in p["trades"]}, key=str.lower), True
    )
    result.created_trades = created_trades
    existing = {
        t.trainee_code.upper(): t
        for t in db.scalars(select(M.Trainee).where(M.Trainee.cycle_id == cycle.id).with_for_update(of=M.Trainee))
    }
    # One lookup + one batched insert for every skill in the file (no per-row SELECT/autoflush).
    row_skills = [clean_names(p["skills"], 20, "skills") for _, p in parsed]
    found, created_skills = resolve_skills_bulk(db, [n for names in row_skills for n in names])
    actives = {
        a.trainee_id: a
        for a in db.scalars(
            select(M.Allocation).where(
                M.Allocation.cycle_id == cycle.id, M.Allocation.status == AllocationStatus.active
            )
        ).unique()
    }
    for (_, p), names in zip(parsed, row_skills, strict=True):
        skill_rows = pick_skills(found, names)
        row_trades = [trades[n.lower()] for n in p["trades"]]
        values = {
            "full_name": p["full_name"],
            "gender": p["gender"],
            "centre_name": p["centre_name"],
            "theory_score": p["theory_score"],
            "practical_score": p["practical_score"],
            "safety_score": p["safety_score"],
            "career_interest": p["career_interest"],
        }
        if p["has_background"]:
            values["background"] = p["background"]
        t = existing.get(p["code"].upper())
        if t is None:
            t = M.Trainee(cycle_id=cycle.id, trainee_code=p["code"], **values)
            set_trainee_trades(t, row_trades)
            t.skills = skill_rows
            db.add(t)
            result.created += 1
            continue
        dirty = (
            any(getattr(t, k) != v for k, v in values.items())
            or sorted(s.id for s in t.skills) != sorted(s.id for s in skill_rows)
            or not same_trades(t.trade_id, t.all_trades, row_trades)
        )
        if not dirty:
            result.unchanged += 1
            continue
        for k, v in values.items():
            setattr(t, k, v)
        set_trainee_trades(t, row_trades)
        t.skills = skill_rows
        result.updated += 1
        a = actives.get(t.id)  # in-memory state already holds the new trades/skills/scores
        if a is not None:
            reasons = eligibility(trainee_input(t), project_input(a.project))
            if reasons:
                result.warnings.append(
                    f"{t.trainee_code} is now ineligible for allocated project "
                    f"{a.project.project_code} ({', '.join(reasons)})"
                )
    db.flush()
    result.created_skills = created_skills
    jobs.enqueue(
        db,
        "allocate_cycle",
        {"cycle_id": str(cycle.id), "trigger": "import"},
        dedupe_key=f"allocate:{cycle.id}",
        trigger="import",
    )
    return Envelope(data=_finish(db, result, dry_run, admin, "import.trainees", cycle))


# ---------------------------------------------------------------- IMP-02 projects

PROJECT_REQUIRED = [
    "Project Title",
    "Target Trade",
    "Difficulty Level",
    "Required Prerequisite Skills",
    "Min Practical Score Required (%)",
    "Project Category / Sponsor Goal",
    "Problem Statement Doc",
]
# "Project ID" is optional (generated when blank; dataset codes are preserved). "Max Trainee Capacity" is a
# legacy column: recognised and kept as history only, never enforced (ADR-017: no team capacity).


@router.post("/admin/imports/projects")
async def import_projects(
    admin: Admin,
    db: DB,
    file: Annotated[UploadFile, File()],
    cycle_id: Annotated[str, Form()],
    dry_run: Annotated[bool, Form()] = True,
) -> Envelope[ImportResult]:
    data = await _read(file)
    sheet = parse_upload(file.filename or "", data, "Project Catalog", "Project Title", PROJECT_REQUIRED)
    cycle = _cycle(db, cycle_id)
    rc = RowCollector(sheet)
    result = ImportResult(
        applied=False,
        rows=len(sheet.rows),
        created=0,
        updated=0,
        unchanged=0,
        created_trades=[],
        created_skills=[],
        warnings=[],
        errors=[],
    )
    if rc.errors:
        result.errors = rc.errors
        return Envelope(data=_finish(db, result, dry_run, admin, "import.projects", cycle))

    has_desc, has_sup = "description" in sheet.headers, "supervisor email" in sheet.headers
    sup_emails = {(_text(r.get("supervisor email")) or "").lower() for _, r in sheet.rows} - {""}
    supervisors = {
        u.email: u
        for u in db.scalars(
            select(M.User).where(
                M.User.email.in_(sup_emails), M.User.role == UserRole.supervisor, M.User.status == UserStatus.active
            )
        )
    }
    existing = {
        p.project_code.upper(): p
        for p in db.scalars(select(M.Project).where(M.Project.cycle_id == cycle.id).with_for_update(of=M.Project))
    }
    alloc_counts = dict(
        db.execute(
            select(M.Allocation.project_id, func.count())
            .where(M.Allocation.cycle_id == cycle.id, M.Allocation.status == AllocationStatus.active)
            .group_by(M.Allocation.project_id)
        ).all()
    )
    titles = {p.title.strip().lower(): p for p in existing.values()}
    seen: set[str] = set()
    parsed: list[dict[str, Any]] = []
    for number, r in sheet.rows:
        before = len(rc.errors)
        code = _text(r.get("project id")) if "project id" in sheet.headers else None
        title = _text(r.get("project title"))
        if not code and title and title.strip().lower() in titles:
            code = titles[title.strip().lower()].project_code  # reimport without codes: match by title
        if code and not CODE_RE.match(code):
            rc.err(number, "Project ID", "letters, digits and '-' only (max 40); leave blank to generate")
        elif code and code.upper() in seen:
            rc.err(number, "Project ID", f"duplicate code {code} in this file")
        elif code:
            seen.add(code.upper())
        if not title or not 5 <= len(title) <= 200:
            rc.err(number, "Project Title", "required (5–200 characters)")
        trades_raw = _split(r.get("target trade"))
        if not trades_raw or len(trades_raw) > 5 or any(len(x) > 80 for x in trades_raw):
            rc.err(number, "Target Trade", "required: 1–5 trades separated by commas (each up to 80 characters)")
        diff_raw = (_text(r.get("difficulty level")) or "").lower()
        if diff_raw not in {d.value for d in Difficulty}:
            rc.err(number, "Difficulty Level", "must be Beginner, Intermediate or Advanced")
        try:
            minimum = _decimal(r.get("min practical score required (%)"))
            if minimum is None or not 0 <= minimum <= 100:
                raise ValueError("required, 0–100")
        except ValueError as e:
            rc.err(number, "Min Practical Score Required (%)", str(e))
            minimum = None
        capacity = None  # legacy column: history only, never enforced
        try:
            cap_d = _decimal(r.get("max trainee capacity")) if "max trainee capacity" in sheet.headers else None
            if cap_d is not None and cap_d == cap_d.to_integral_value() and 1 <= cap_d <= 100:
                capacity = int(cap_d)
        except ValueError:
            capacity = None
        skills = _split(r.get("required prerequisite skills"))
        if len(skills) > 15 or any(len(s) > 80 for s in skills):
            rc.err(number, "Required Prerequisite Skills", "at most 15 skills, each up to 80 characters")
        category = _text(r.get("project category / sponsor goal"))
        if category and len(category) > 120:
            rc.err(number, "Project Category / Sponsor Goal", "at most 120 characters")
        doc = _text(r.get("problem statement doc"))
        if doc and (not doc.startswith("https://") or len(doc) > 2048):
            rc.err(number, "Problem Statement Doc", "must be an https:// link")
        description = _text(r.get("description")) if has_desc else None
        if description and len(description) > 5000:
            rc.err(number, "Description", "at most 5000 characters")
        sup_email = (_text(r.get("supervisor email")) or "").lower() if has_sup else ""
        supervisor = supervisors.get(sup_email) if sup_email else None
        if sup_email and supervisor is None:
            rc.err(number, "Supervisor Email", "must belong to an active supervisor")
        current = existing.get(code.upper()) if code else None
        if current is not None and has_sup and not sup_email and current.status == ProjectStatus.published:
            rc.err(number, "Supervisor Email", "a published project needs a supervisor")
        if len(rc.errors) == before:
            parsed.append(
                {
                    "code": code,
                    "title": title,
                    "trades": clean_names(trades_raw, 5, "Target Trade") if trades_raw else [],
                    "difficulty": Difficulty(diff_raw),
                    "min_practical_score": minimum,
                    "max_trainees": capacity,
                    "skills": skills,
                    "category": category,
                    "problem_doc_url": doc,
                    "description": description,
                    "supervisor": supervisor,
                }
            )
    result.errors = rc.errors
    if result.errors:
        return Envelope(data=_finish(db, result, dry_run, admin, "import.projects", cycle))

    trades, result.created_trades = resolve_trade_names(
        db, sorted({n for p in parsed for n in p["trades"]}, key=str.lower), True
    )
    from app.modules.projects import next_project_code

    touched: list[M.Project] = []
    row_skills = [clean_names(p["skills"], 15, "skills") for p in parsed]
    found, result.created_skills = resolve_skills_bulk(db, [n for names in row_skills for n in names])
    for p, names in zip(parsed, row_skills, strict=True):
        skill_rows = pick_skills(found, names)
        row_trades = [trades[n.lower()] for n in p["trades"]]
        values: dict[str, Any] = {
            "title": p["title"],
            "difficulty": p["difficulty"],
            "min_practical_score": p["min_practical_score"],
            "max_trainees": p["max_trainees"],
            "category": p["category"],
            "problem_doc_url": p["problem_doc_url"],
        }
        if has_desc:
            values["description"] = p["description"]
        if has_sup:
            values["supervisor_id"] = p["supervisor"].id if p["supervisor"] else None
        proj = existing.get(p["code"].upper()) if p["code"] else None
        if proj is None:
            proj = M.Project(
                id=uuid.uuid4(),
                cycle_id=cycle.id,
                project_code=p["code"] or next_project_code(db, cycle.id),
                created_by=admin.id,
                generation_source="import",
                processing_status=ProjectProcessing.queued  # ADR-028: a description is enough to process
                if p["problem_doc_url"] or (values.get("description") or "").strip()
                else ProjectProcessing.awaiting_document,
                **values,
            )
            set_project_trades(proj, row_trades)
            proj.skills = skill_rows
            db.add(proj)
            touched.append(proj)
            result.created += 1
            continue
        dirty = (
            any(getattr(proj, k) != v for k, v in values.items() if k != "max_trainees")
            or sorted(s.id for s in proj.skills) != sorted(s.id for s in skill_rows)
            or not same_trades(proj.target_trade_id, proj.all_trades, row_trades)
        )
        if not dirty:
            result.unchanged += 1
            continue
        trade_or_min = (
            not same_trades(proj.target_trade_id, proj.all_trades, row_trades)
            or proj.min_practical_score != values["min_practical_score"]
        )
        for k, v in values.items():
            setattr(proj, k, v)
        set_project_trades(proj, row_trades)
        proj.skills = skill_rows
        if proj.processing_status == ProjectProcessing.awaiting_document and (
            p["problem_doc_url"] or (proj.description or "").strip()
        ):
            proj.processing_status = ProjectProcessing.queued
        touched.append(proj)
        result.updated += 1
        if trade_or_min and alloc_counts.get(proj.id):
            db.flush()
            db.refresh(proj)
            pi = project_input(proj)
            for a in db.scalars(
                select(M.Allocation).where(
                    M.Allocation.project_id == proj.id, M.Allocation.status == AllocationStatus.active
                )
            ).unique():
                if reasons := eligibility(trainee_input(a.trainee), pi):
                    result.warnings.append(
                        f"{a.trainee.trainee_code} is now ineligible for {proj.project_code} ({', '.join(reasons)})"
                    )
    db.flush()
    for i in range(0, len(touched), 200):  # one batch job per 200 projects instead of one job each
        ids = [str(proj.id) for proj in touched[i : i + 200]]
        jobs.enqueue(db, "process_projects", {"project_ids": ids}, trigger="import")
    return Envelope(data=_finish(db, result, dry_run, admin, "import.projects", cycle))
